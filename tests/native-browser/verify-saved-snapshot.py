"""Independent downloaded saved-snapshot reader; never contacts a provider."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import zipfile


def check(directory):
    root = Path(directory)
    doc = json.loads((root / 'snapshot.json').read_text(encoding='utf-8'))
    assert doc['scope']['kind'] == 'saved_snapshot'
    records = doc['records']
    assert 1 <= len(records) <= 100
    assert doc['counts']['exportedRecords'] == doc['counts']['scopeRecords'] == len(records)
    assert doc['counts']['providerMatches'] is None
    assert doc['counts']['retrievedRecords'] is None
    ids = [r['searchId'] for r in records]
    assert len(set(ids)) == len(ids)
    queries = {q['runId'] for q in doc['queryContexts']}
    for r in records:
        assert r['runIds'] and set(r['runIds']) <= queries
        for value in r['identifiers'].values():
            assert value is None or isinstance(value, str)
        assert not ({'password', 'password_hash', 'csrf', 'token', 'database'} & set(r))
    lines = [json.loads(line) for line in (root / 'snapshot.jsonl').read_text(encoding='utf-8').splitlines()]
    # Read the documented NDJSON header plus individual record envelopes.
    observed = [line['record'] for line in lines if line.get('type') == 'record']
    assert observed == records
    rows = list(csv.DictReader(io.StringIO((root / 'snapshot.csv').read_text(encoding='utf-8'))))
    assert [r['Search ID'] for r in rows] == ids
    safe = lambda value: "'" + value if value.lstrip(' \t\r\n')[:1] in '=+-@' and value.lstrip(' \t\r\n') else value
    for row, record in zip(rows, records):
        for key, field in [('PMID', 'pmid'), ('PMCID', 'pmcid'), ('DOI', 'doi')]:
            assert row[key] == safe(record['identifiers'][field] or '')
        assert row['Search Run ID'] == '; '.join(record['runIds'])
        assert row['Reason'] == safe(record['acquisition']['reason'] or '')
        assert row['Abstract'] == safe(record['publication']['abstract'] or '')
        assert json.loads(row['Typed record JSON']) == record
        assert json.loads(row['Original provenance JSON']) == record['originals']
        assert {q['runId'] for q in json.loads(row['Query contexts JSON'])} == set(record['runIds'])
    from openpyxl import load_workbook
    book = load_workbook(root / 'snapshot.xlsx', read_only=False, data_only=False)
    sheet = book['Literature']
    assert sheet.max_row == len(records) + 1
    for rank, record in enumerate(records, 2):
        for col, expected in [(1, record['searchId']), (5, record['identifiers']['pmid']), (6, record['identifiers']['pmcid']), (7, record['identifiers']['doi'])]:
            cell = sheet.cell(rank, col)
            assert (cell.value or '') == (expected or '') and cell.data_type != 'f'
        assert sheet.cell(rank, 20).value == '; '.join(record['runIds'])
        assert (sheet.cell(rank, 23).value or '') == (record['publication']['abstract'] or '')
        assert json.loads(sheet.cell(rank, 28).value) == record
    with zipfile.ZipFile(root / 'snapshot.zip') as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == len(set(archive.namelist()))
        manifest = json.loads(archive.read('manifest.json'))
        assert manifest['plan']['scopeKind'] == manifest['plan']['state'] == 'saved_snapshot'
        assert manifest['plan']['allowedActions'] == []
        assert [r['searchId'] for r in manifest['items']] == ids
        files, included = set(), 0
        for item in manifest['items']:
            assert item['childBatchId'] is None
            if item['availability'] == 'included':
                data = archive.read(item['file'])
                assert hashlib.sha256(data).hexdigest() == item['original']['sha256']
                assert len(data) == item['original']['bytes']
                assert data.startswith(b'%PDF-')
                files.add(item['file']); included += 1
            else:
                assert item['availabilityReason']
        assert set(archive.namelist()) == files | {'manifest.json', 'records.json'}
        assert manifest['counts']['members'] == len(records)
        assert manifest['counts']['includedRecords'] == included
        assert manifest['counts']['unresolvedRecords'] == len(records) - included
        assert manifest['counts']['uniqueOriginals'] == len(files)
    return {'records': len(records), 'runAssociations': sum(len(r['runIds']) for r in records), 'uniqueRuns': len(queries), 'includedOriginalAssociations': included, 'uniqueFiles': len(files), 'formats': ['JSON', 'JSONL', 'CSV', 'XLSX', 'ZIP'], 'physicalExcelOrDeviceQualification': False}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('directory')
    args = parser.parse_args()
    print(json.dumps(check(args.directory), ensure_ascii=False))
