"""Independent XLSX acceptance reader; never contacts a source provider.

Expected cells/hyperlinks are supplied by the caller from independent API or
fixture inputs. Negative controls deliberately change the emitted package.
"""
import argparse
import hashlib
import io
import json
import pathlib
import sys
import tempfile
import zipfile
from urllib.parse import urlsplit

from defusedxml import ElementTree as ET
import openpyxl

if sys.flags.optimize:
    raise SystemExit('Optimized Python disables acceptance assertions and is unsupported.')

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
HEADERS = ['Search ID', 'Title', 'Authors', 'Year', 'PMID', 'PMCID', 'DOI',
           'PubMed URL', 'PMC URL', 'DOI URL', 'Current item state', 'Reason',
           'Rights URI', 'Original SHA256', 'Acquisition URL', 'Repository datestamp',
           'Original format', 'Original version', 'Original bytes', 'Search Run ID',
           'Batch ID', 'DOI link status']


def verify(path, expected):
    raw = pathlib.Path(path).read_bytes()
    assert 0 < len(raw) <= 8 * 1024 * 1024, 'workbook file bound'
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = archive.infolist()
        assert len(members) == len({x.filename for x in members}) <= 64, 'duplicate/excess package members'
        assert sum(x.file_size for x in members) <= 32 * 1024 * 1024, 'expanded package bound'
        assert archive.testzip() is None, 'ZIP CRC'
        for member in members:
            assert not member.filename.startswith('/') and '..' not in pathlib.PurePosixPath(member.filename).parts
            assert 'vba' not in member.filename.lower() and 'externalLink' not in member.filename, 'active workbook content'
            if member.filename.endswith(('.xml', '.rels')):
                root = ET.fromstring(archive.read(member))
                assert not list(root.iter(NS + 'f')), 'formula node'
                for relation in root:
                    if relation.get('TargetMode') == 'External':
                        assert relation.get('Type', '').endswith('/hyperlink'), 'external data relationship'
                        link = urlsplit(relation.get('Target', ''))
                        assert link.scheme == 'https' and link.hostname in {
                            'doi.org', 'pmc.ncbi.nlm.nih.gov', 'pubmed.ncbi.nlm.nih.gov', 'creativecommons.org'
                        } and not link.username and not link.password and not link.port and not link.fragment, 'unsafe hyperlink'
                if member.filename.startswith('xl/worksheets/') and member.filename.endswith('.xml'):
                    assert all(c.get('t') in ('s', 'inlineStr') for c in root.iter(NS + 'c')), 'non-text XML cell'
    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    assert workbook.sheetnames == ['Literature'], 'sheet identity'
    sheet = workbook['Literature']
    assert sheet.max_row == expected['rows'] and sheet.max_column == len(HEADERS), 'dimensions'
    assert [sheet.cell(1, i + 1).value for i in range(len(HEADERS))] == HEADERS, 'column contract'
    for row in sheet:
        for cell in row:
            assert cell.data_type == 's' and cell.number_format == '@', 'reader cell type/format'
    for coordinate, value in expected['cells'].items():
        assert (sheet[coordinate].value or '') == value, 'cell mismatch: ' + coordinate
    actual_links = {cell.coordinate: cell.hyperlink.target for row in sheet for cell in row if cell.hyperlink}
    assert actual_links == expected['hyperlinks'], 'hyperlink set/target mismatch'
    workbook.close()
    return {'sha256': hashlib.sha256(raw).hexdigest(), 'rows': sheet.max_row - 1,
            'columns': len(HEADERS), 'hyperlinks': len(actual_links), 'reader': openpyxl.__version__}


def synthetic_expected():
    # Literal expectations are independent of the Go encoder and its link builder.
    return {'rows': 3, 'cells': {
        'A2': '000000000000000000000001', 'B2': '=HYPERLINK("https://untrusted.invalid","SYNTHETIC")',
        'C2': '王小明 · α · 😀 · é\nSecond line', 'D2': '002026', 'E2': '0012345678901234567890',
        'F2': 'PMC000001', 'G2': '10.0000/a?b#c%25', 'K2': 'acquired',
        'N2': 'a' * 64, 'P2': '2026-01-02', 'Q2': 'XML',
        'R2': 'repository snapshot; publication version unspecified', 'S2': '123', 'T2': 'RUN-SYNTHETIC-0001',
        'B3': '+SUM(1,1)', 'C3': '-1+2', 'E3': '00002', 'K3': 'unavailable',
        'L3': '@SUM(1,1) — Unknown rights; open source links.\r\nLiteral _x0041_ stays literal.',
        'M3': 'javascript:alert(1)', 'O3': 'https://user:secret@pmc.ncbi.nlm.nih.gov/private',
        'S3': '0', 'U3': '',
    }, 'hyperlinks': {
        'H2': 'https://pubmed.ncbi.nlm.nih.gov/0012345678901234567890/',
        'I2': 'https://pmc.ncbi.nlm.nih.gov/articles/PMC000001/',
        'J2': 'https://doi.org/10.0000/a%3Fb%23c%2525',
        'M2': 'https://creativecommons.org/licenses/by/4.0/',
        'O2': 'https://pmc.ncbi.nlm.nih.gov/api/oai/v1/mh/?verb=GetRecord&identifier=oai%3Apubmedcentral.nih.gov%3A000001',
        'H3': 'https://pubmed.ncbi.nlm.nih.gov/00002/', 'I3': 'https://pmc.ncbi.nlm.nih.gov/articles/PMC000002/',
    }}


def controls(path, expected):
    with zipfile.ZipFile(path) as z:
        source = {i.filename: z.read(i) for i in z.infolist()}
    rejected = []
    for name in ('numeric-cell', 'formula-cell', 'identifier-change', 'unsafe-link', 'missing-row'):
        changed = dict(source)
        sheet = ET.fromstring(changed['xl/worksheets/sheet1.xml'])
        cell = next(c for c in sheet.iter(NS + 'c') if c.get('r') == 'A2')
        if name == 'numeric-cell':
            cell.set('t', 'n')
        elif name == 'formula-cell':
            import xml.etree.ElementTree as stdET
            stdET.SubElement(cell, NS + 'f').text = '1+1'
        elif name == 'identifier-change':
            strings = ET.fromstring(changed['xl/sharedStrings.xml'])
            next(n for n in strings.iter(NS + 't') if n.text == expected['cells']['A2']).text = '1'
            import xml.etree.ElementTree as stdET
            changed['xl/sharedStrings.xml'] = stdET.tostring(strings)
        elif name == 'unsafe-link':
            key = 'xl/worksheets/_rels/sheet1.xml.rels'
            relations = ET.fromstring(changed[key])
            relations[0].set('Target', 'file:///private')
            import xml.etree.ElementTree as stdET
            changed[key] = stdET.tostring(relations)
        elif name == 'missing-row':
            data = sheet.find(NS + 'sheetData')
            data.remove(data[-1])
        import xml.etree.ElementTree as stdET
        changed['xl/worksheets/sheet1.xml'] = stdET.tostring(sheet)
        with tempfile.TemporaryDirectory() as temporary:
            target = pathlib.Path(temporary) / 'mutated.xlsx'
            with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as z:
                for member, data in changed.items():
                    z.writestr(member, data)
            try:
                verify(target, expected)
            except (AssertionError, ValueError, KeyError):
                rejected.append(name)
            else:
                raise AssertionError('mutation survived: ' + name)
    return rejected


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=pathlib.Path)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--synthetic', action='store_true')
    group.add_argument('--expected', type=pathlib.Path)
    parser.add_argument('--negative-controls', action='store_true')
    args = parser.parse_args()
    expectations = synthetic_expected() if args.synthetic else json.loads(args.expected.read_text(encoding='utf-8'))
    result = verify(args.workbook, expectations)
    if args.negative_controls:
        assert args.synthetic, 'negative controls require explicit synthetic fixture'
        result['rejected_mutations'] = controls(args.workbook, expectations)
    result['scope'] = 'synthetic fixture' if args.synthetic else 'independent supplied expectations'
    print(json.dumps(result, sort_keys=True))
